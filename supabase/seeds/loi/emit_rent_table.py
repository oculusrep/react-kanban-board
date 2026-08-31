#!/usr/bin/env python3
"""LOI rent-table emitter (SPIKE, end-cap drive-thru only).

Creates the 11x4 rent table (Years / Monthly / Yearly / Per Square Foot) at the R1 insertion point
in the blank template, populated with a computed schedule. Structure/format matched to the worked
ECDT deliverable (11x4, "Normal Table" style, no borders, col widths [2080,1380,1500,1660] twips,
blank spacer row, "Extension Options:" label row).

Table CREATION at a known insertion point (template has 0 tables), not editing. The assembler does no
math — rows come from the OVIS rent engine. Freestanding is a DIFFERENT shape and out of scope here.
"""
import sys
import docx
from docx.oxml.ns import qn
from decimal import Decimal
from rent_engine import Terms, compute_schedule

COLW = [2080, 1380, 1500, 1660]  # twips, from worked ECDT tblGrid


def money(x: Decimal) -> str:
    return f"${x:,.2f}"


def logical_rows(sched, base_term_years: int):
    """11-row logical model: header, base-term rows, blank spacer, 'Extension Options:', option rows."""
    base = [r for r in sched if r.end_year <= base_term_years]
    opts = [r for r in sched if r.start_year > base_term_years]
    def cells(r):
        return [f"{r.start_year}-{r.end_year}", money(r.monthly), money(r.yearly), money(r.psf)]
    rows = [["Years", "Monthly", "Yearly", "Per Square Foot"]]
    rows += [cells(r) for r in base]
    rows.append(["", "", "", ""])                       # blank spacer
    rows.append(["Extension Options:", "", "", ""])     # label row
    rows += [cells(r) for r in opts]
    return rows


def _set_col_widths(table, widths):
    # tblGrid
    grid = table._tbl.find(qn('w:tblGrid'))
    for gc, w in zip(grid.findall(qn('w:gridCol')), widths):
        gc.set(qn('w:w'), str(w))
    # per-cell width (Word honors these)
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn('w:tcW'))
            if tcW is None:
                tcW = tcPr.makeelement(qn('w:tcW'), {})
                tcPr.append(tcW)
            tcW.set(qn('w:w'), str(w)); tcW.set(qn('w:type'), 'dxa')


def emit(template_path, out_path, terms, base_term_years):
    doc = docx.Document(template_path)
    paras = doc.paragraphs
    anchor = next(p for p in paras if "TIED TO SQUARE FOOTAGE" in p.text)
    note = next(p for p in paras if "based on Landlord" in p.text)

    # remove the tab-delimited R1 block strictly between the anchor instruction and the note
    start_i = paras.index(anchor) + 1
    end_i = paras.index(note)
    for p in paras[start_i:end_i]:
        p._p.getparent().remove(p._p)

    sched = compute_schedule(terms)
    data = logical_rows(sched, base_term_years)
    table = doc.add_table(rows=len(data), cols=4)
    table.style = "Normal Table"
    table.autofit = False
    for ri, row in enumerate(data):
        for ci, val in enumerate(row):
            table.rows[ri].cells[ci].text = val
    _set_col_widths(table, COLW)

    # move the appended table to the insertion point (right after the anchor instruction)
    anchor._p.addnext(table._tbl)
    doc.save(out_path)
    return table


def diff_against(out_path, ref_path):
    """Report differences vs the worked ECDT rent table — these are bugs, not cosmetics."""
    out = docx.Document(out_path)
    ref = docx.Document(ref_path)
    diffs = []
    ot = next((t for t in out.tables if len(t.columns) == 4), None)
    rt = ref.tables[0]
    if ot is None:
        return ["no 4-col table emitted"]
    if (len(ot.rows), len(ot.columns)) != (len(rt.rows), len(rt.columns)):
        diffs.append(f"dims {len(ot.rows)}x{len(ot.columns)} vs ref {len(rt.rows)}x{len(rt.columns)}")
    for ri in range(min(len(ot.rows), len(rt.rows))):
        for ci in range(4):
            a, b = ot.rows[ri].cells[ci].text, rt.rows[ri].cells[ci].text
            if a != b:
                diffs.append(f"cell r{ri}c{ci}: {a!r} vs ref {b!r}")
    og = [g.get(qn('w:w')) for g in ot._tbl.find(qn('w:tblGrid')).findall(qn('w:gridCol'))]
    rg = [g.get(qn('w:w')) for g in rt._tbl.find(qn('w:tblGrid')).findall(qn('w:gridCol'))]
    if og != rg:
        diffs.append(f"col widths {og} vs ref {rg}")
    if (ot.style.name if ot.style else None) != (rt.style.name if rt.style else None):
        diffs.append(f"style {ot.style.name} vs ref {rt.style.name}")
    return diffs


if __name__ == "__main__":
    D = Decimal
    base = "/Users/mike/Documents/GitHub/react-kanban-board-loi/supabase/seeds/loi/"
    tmpl = base + "templates/LOI_US_7_30_2026.docx"
    out = base + "fixtures/SPIKE_powder_springs_emitted.docx"
    ref = base + "fixtures/LOI US - ECDT (072025).orep.docx"
    ps = Terms("per_sqft", D("60.00"), D("0.10"), 5, 40, sqft=2109)
    emit(tmpl, out, ps, base_term_years=10)
    print("emitted:", out)
    diffs = diff_against(out, ref)
    if not diffs:
        print("DIFF vs worked ECDT: NONE — table matches (cells, dims, col widths, style).")
    else:
        print(f"DIFF vs worked ECDT: {len(diffs)} difference(s):")
        for d in diffs:
            print("  -", d)
